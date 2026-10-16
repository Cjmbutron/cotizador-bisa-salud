import openpyxl, json, math, warnings
from decimal import Decimal as D, ROUND_HALF_UP, ROUND_CEILING
from openpyxl.utils import column_index_from_string as ci
warnings.filterwarnings('ignore')
import os
SRC=os.environ.get('BISA_XLSM','/tmp/ref.xlsm')
wb=openpyxl.load_workbook(SRC,data_only=True)
IGR=wb['PRIMAS IGR RMA ADV25']; MBL=wb['PRIMAS MBL']; AR=wb['PRIMAS ALTO RIESGO']
REG={'LA PAZ':'B','COCHABAMBA':'M','SANTA CRUZ':'X','CHUQUISACA':'AI','TARIJA':'AT'}
def num(v): return v if isinstance(v,(int,float)) and not isinstance(v,bool) else None
def rhu(x): return int(D(str(x)).quantize(D('1'),rounding=ROUND_HALF_UP))   # REDONDEAR
def ceil(x): return int(D(str(x)).quantize(D('1'),rounding=ROUND_CEILING))  # REDONDEAR.MAS / entero superior
PARAM=dict(SARS_IG=10,SARS_RM=70,MAT_IG=20,MAT_RM=140,RM_DESC={'700':0,'140':28,'70':49},
           CONTADO={'IG':D('0.975'),'RM':D('0.975'),'AD':D('0.975'),'MB':D('0.909'),'NP':1/D('1.1')},
           MB_EDAD_ADULTO={'TITULAR':18,'CONYUGE':18,'HIJ@':25})
def rate_igr(prod,region,sexcode,age):
    c=ci(REG[region])+{'IG':0,'RM':3,'AD':6}[prod]+{'M':0,'F-M':1,'F':2}[sexcode]
    return num(IGR.cell(4+age,c).value)   # None => N/A
COV={'1MM':0,'2MM':6,'5MM':12}; DED={'1K':0,'2K':1,'3K':2,'5K':3,'10K':4,'20K':5}
def rate_mb_adult(cov,ded,age):
    if age<18: return None
    return num(MBL.cell(6+age-18,3+COV[cov]+DED[ded]).value)
def rate_mb_child(cov,ded,n):
    for r in range(6,13):
        if MBL.cell(r,22).value==n: return num(MBL.cell(r,24+COV[cov]+DED[ded]).value)
    return None
VAI={'50K':0,'100K':1,'150K':2,'250K':3,'500K':4,'1MM':5,'2MM':6}; SC={'L':0,'U':1,'M':2,'MB':3}
DDI={'2K':0,'3K':1,'4K':2,'5K':3,'10K':4,'15K':5,'20K':6,'25K':7,'30K':8,'40K':9,'50K':10,'100K':11}
def rate_np(scope,va,ded,age):
    c=2+12*VAI[va]+84*SC[scope]+DDI[ded]; return num(AR.cell(18+age,c).value)
def np_monthly(annual): return rhu(D(str(rhu(D(ceil(D(annual)/12))/D('1.2675')*D('1.3'))))*D('1.1'))
def quote(q):
    out=[]
    for p in q['people']:
        rel,age,sex=p['rel'],p['age'],p['sex']; code={'MASCULINO':'M','FEMENINO SIN MATERNIDAD':'F','FEMENINO CON MATERNIDAD':'F-M'}.get(sex,0)
        row={}
        # IG/RM/AD
        for prod in ('IG','RM','AD'):
            if not q['on'].get(prod) or code==0: row[prod]=0; continue
            mat=(code=='F-M')
            if mat and not (19<=age<=49): row[prod]='N/A'; continue
            base=rate_igr(prod,q['region'],'F' if (mat and prod in('IG','RM')) else code,age)
            if base is None and not (mat and prod=='AD'): row[prod]='N/A'; continue
            if base is None: row[prod]='N/A'; continue
            if prod=='RM':
                base=base-PARAM['RM_DESC'][q['rm']]
                if base<1: base=0
            tot=base; sars=0
            if mat and prod=='IG': tot+=PARAM['MAT_IG']
            if mat and prod=='RM': tot+=PARAM['MAT_RM']
            row[prod+'_nosars']=tot
            if q['sars'] and tot>0: tot+= {'IG':10,'RM':70,'AD':0}[prod]
            row[prod]=tot
        # MB
        if not q['on'].get('MB'): row['MB']=0
        else:
            thr=PARAM['MB_EDAD_ADULTO']['TITULAR' if rel in('TITULAR','CONYUGE') else 'HIJ@']
            if age>=thr:
                v=rate_mb_adult(q['mb_cov'],q['mb_ded'],age); row['MB']='N/A' if v is None else v
            elif not rel: row['MB']=0
            else:
                v=rate_mb_child(q['mb_cov'],q['mb_ded'],q['mb_n'])
                row['MB']='ERR' if v is None else ceil(v)
        # NP
        if age>64 or (1<=age<=18 and code=='F-M'): row['NP']='N/A'
        elif age<=0: row['NP']=0
        else:
            a=rate_np(*q['np'],age); row['NP']='N/A' if a is None else np_monthly(a)
        out.append(row)
    return out
