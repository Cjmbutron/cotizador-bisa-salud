#!/usr/bin/env python3
"""
Convierte las 13 hojas de propuesta del Excel (Infinity CM/SM, RedMax CM/SM, Advance CM/SM,
MB 1/2/5 CM/SM, Naranja +) en plantillas de datos (JSON + imágenes) que el servidor dibuja como PDF.

Uso:  python3 -I extraer_plantillas.py <archivo.xlsm> <carpeta_salida>

Cada celda conserva texto, estilo y combinación. Las fórmulas se guardan como árbol ("f") para que el
servidor las evalúe con los datos de la cotización; el valor que Excel dejó guardado queda en "v" y sirve
para verificar que el evaluador reproduce al Excel.
"""
import json, os, re, sys, warnings, colorsys, hashlib
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter, column_index_from_string

HOJAS = {
    'Infinity CM': 'IG_CM', 'Infinity SM': 'IG_SM',
    'RedMax CM': 'RM_CM', 'RedMax SM': 'RM_SM',
    'Advance CM': 'AD_CM', 'Advance SM': 'AD_SM',
    'MB 1 CM': 'MB_1MM_CM', 'MB 1 SM': 'MB_1MM_SM',
    'MB 2 CM': 'MB_2MM_CM', 'MB 2 SM': 'MB_2MM_SM',
    'MB 5 CM': 'MB_5MM_CM', 'MB 5 SM': 'MB_5MM_SM',
    'Naranja +': 'NP',
}

def px_ancho(w):  # ancho de columna Excel (caracteres) -> píxeles
    return round((w if w is not None else 8.43) * 7 + 5)

def pt_a_px(pt):
    return pt * 96 / 72

TEMA = None
def cargar_tema(wb):
    global TEMA
    xml = wb.loaded_theme.decode('utf8', 'ignore') if wb.loaded_theme else ''
    orden = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6']
    TEMA = []
    for k in orden:
        m = re.search(r'<a:%s>.*?(?:val|lastClr)="([0-9A-Fa-f]{6})"' % k, xml, re.S)
        TEMA.append(m.group(1) if m else '000000')

INDEXADO = {0: '000000', 1: 'FFFFFF', 2: 'FF0000', 3: '00FF00', 4: '0000FF', 5: 'FFFF00', 6: 'FF00FF', 7: '00FFFF',
            8: '000000', 9: 'FFFFFF'}

def color(c):
    if c is None:
        return None
    try:
        if c.type == 'rgb' and isinstance(c.rgb, str) and len(c.rgb) in (6, 8):
            h = c.rgb[-6:]
            return None if c.rgb.startswith('00') and len(c.rgb) == 8 and h == '000000' and False else '#' + h
        if c.type == 'theme':
            h = TEMA[c.theme]
            if c.tint:
                r, g, b = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
                hh, l, s = colorsys.rgb_to_hls(r, g, b)
                l = l * (1 + c.tint) if c.tint < 0 else l * (1 - c.tint) + c.tint
                r, g, b = colorsys.hls_to_rgb(hh, max(0, min(1, l)), s)
                h = '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))
            return '#' + h
        if c.type == 'indexed':
            return '#' + INDEXADO.get(c.indexed, '000000')
    except Exception:
        return None
    return None

BORDE = {'thin': '1px solid', 'medium': '2px solid', 'thick': '3px solid', 'hair': '0.5px solid', 'dotted': '1px dotted',
         'dashed': '1px dashed', 'double': '3px double', 'mediumDashed': '2px dashed'}

def borde(b):
    if b is None or b.style is None:
        return None
    return '%s %s' % (BORDE.get(b.style, '1px solid'), ('#9a9a9a' if b.style == 'hair' else (color(b.color) or '#000')))

def clase_formato(fmt):
    if fmt is None or fmt == 'General':
        return 'g'
    if '%' in fmt:
        return 'pct'
    if 'yyyy' in fmt:
        return 'fecha'
    if '"-"' in fmt:
        return 'ac'
    if '0.00' in fmt:
        return 'n2'
    if fmt == '0':
        return 'n0'
    return 'g'


TOKEN = re.compile(r"""\s*(?:
    (?P<str>"(?:[^"]|"")*")
  | (?P<err>(?:(?:\[\d+\])?COTIZADOR!)?\#REF!)
  | (?P<xref>(?:\[\d+\])?COTIZADOR!\$?[A-Z]{1,3}\$?\d+)
  | (?P<ref>\$?[A-Z]{1,3}\$?\d+)(?![A-Za-z(])
  | (?P<num>\d+(?:\.\d+)?)
  | (?P<fn>[A-Z]+)\(
  | (?P<op><>|<=|>=|[-+*/&=<>(),])
)""", re.X)

class Analizador:
    """Analiza el subconjunto de Excel usado en las plantillas: IF, OR, TODAY, + - * / & y comparaciones."""
    def __init__(self, texto):
        self.t = []
        pos = 0
        while pos < len(texto):
            m = TOKEN.match(texto, pos)
            if not m:
                raise ValueError('token inesperado en %r posición %d' % (texto, pos))
            pos = m.end()
            k = m.lastgroup
            self.t.append((k, m.group(k)))
        self.i = 0

    def ver(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def tomar(self):
        x = self.ver(); self.i += 1; return x

    def expresion(self):
        return self.comparacion()

    def comparacion(self):
        a = self.concat()
        while self.ver()[0] == 'op' and self.ver()[1] in ('=', '<>', '<', '>', '<=', '>='):
            op = self.tomar()[1]; a = ['cmp', op, a, self.concat()]
        return a

    def concat(self):
        a = self.suma()
        while self.ver() == ('op', '&'):
            self.tomar(); a = ['cat', a, self.suma()]
        return a

    def suma(self):
        a = self.producto()
        while self.ver()[0] == 'op' and self.ver()[1] in '+-' and self.ver()[1] != '':
            op = self.tomar()[1]; a = ['bin', op, a, self.producto()]
        return a

    def producto(self):
        a = self.unario()
        while self.ver()[0] == 'op' and self.ver()[1] in ('*', '/'):
            op = self.tomar()[1]; a = ['bin', op, a, self.unario()]
        return a

    def unario(self):
        if self.ver() == ('op', '+'):
            self.tomar(); return self.unario()
        if self.ver() == ('op', '-'):
            self.tomar(); return ['neg', self.unario()]
        return self.atomo()

    def atomo(self):
        k, v = self.tomar()
        if k == 'str': return ['s', v[1:-1].replace('""', '"')]
        if k == 'num': return ['n', float(v) if '.' in v else int(v)]
        if k == 'err': return ['err']
        if k == 'xref': return ['x', v.split('!')[1].replace('$', '')]
        if k == 'ref': return ['c', v.replace('$', '')]
        if k == 'fn':
            args = []
            if self.ver() != ('op', ')'):
                args.append(self.expresion())
                while self.ver() == ('op', ','):
                    self.tomar(); args.append(self.expresion())
            assert self.tomar() == ('op', ')'), 'falta )'
            return [v.lower()] + args
        if (k, v) == ('op', '('):
            a = self.expresion(); assert self.tomar() == ('op', ')'); return a
        raise ValueError('átomo inesperado %r' % (v,))

def analizar(f, informe, hoja, coord):
    if not isinstance(f, str) or not f.startswith('='):
        return None
    try:
        a = Analizador(f[1:])
        arbol = a.expresion()
        if a.i != len(a.t):
            raise ValueError('sobran tokens')
        return arbol
    except Exception as e:
        informe.append('%s!%s: no se pudo analizar (%s): %s' % (hoja, coord, e, f[:120]))
        return None


NS = {'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
EMU = 9525.0

def _color_dml(el):
    """Color de un nodo DrawingML (srgbClr / schemeClr) como #RRGGBB."""
    if el is None:
        return None
    sc = el.find('a:srgbClr', NS)
    if sc is not None:
        return '#' + sc.get('val')
    sc = el.find('a:schemeClr', NS)
    if sc is not None:
        mapa = {'bg1': 0, 'lt1': 0, 'tx1': 1, 'dk1': 1, 'bg2': 2, 'lt2': 2, 'tx2': 3, 'dk2': 3, 'accent1': 4, 'accent2': 5,
                'accent3': 6, 'accent4': 7, 'accent5': 8, 'accent6': 9}
        i = mapa.get(sc.get('val'))
        return '#' + TEMA[i] if i is not None else None
    return None

def leer_dibujos(ruta):
    """Lee, por hoja, las imágenes (con su recorte) y los cuadros de texto de los dibujos del libro."""
    import zipfile, posixpath
    import xml.etree.ElementTree as ET
    z = zipfile.ZipFile(ruta)
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    wrel = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    dest = {r.get('Id'): r.get('Target') for r in wrel}
    out = {}
    for sh in wb.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheets'):
        nombre = sh.get('name')
        t = dest[sh.get('{%s}id' % NS['r'])].lstrip('/')
        t = t if t.startswith('xl/') else 'xl/' + t
        rels = posixpath.join(posixpath.dirname(t), '_rels', posixpath.basename(t) + '.rels')
        if rels not in z.namelist():
            continue
        dr = [r.get('Target') for r in ET.fromstring(z.read(rels)) if r.get('Type', '').endswith('/drawing')]
        if not dr:
            continue
        dpath = posixpath.normpath(posixpath.join(posixpath.dirname(t), dr[0]))
        drels = posixpath.join(posixpath.dirname(dpath), '_rels', posixpath.basename(dpath) + '.rels')
        medios = {r.get('Id'): posixpath.normpath(posixpath.join(posixpath.dirname(dpath), r.get('Target')))
                  for r in ET.fromstring(z.read(drels))} if drels in z.namelist() else {}
        raiz = ET.fromstring(z.read(dpath))
        items = []
        for anc in raiz:
            tipo = anc.tag.split('}')[1]
            if tipo not in ('twoCellAnchor', 'oneCellAnchor', 'absoluteAnchor'):
                continue
            pos = {'tipo': tipo}
            def punto(n):
                return {'col': int(n.find('xdr:col', NS).text), 'dx': int(n.find('xdr:colOff', NS).text) / EMU,
                        'fila': int(n.find('xdr:row', NS).text), 'dy': int(n.find('xdr:rowOff', NS).text) / EMU}
            if tipo != 'absoluteAnchor':
                pos['desde'] = punto(anc.find('xdr:from', NS))
            if tipo == 'twoCellAnchor':
                pos['hasta'] = punto(anc.find('xdr:to', NS))
            ext = anc.find('xdr:ext', NS)
            if ext is not None:
                pos['ext'] = (int(ext.get('cx')) / EMU, int(ext.get('cy')) / EMU)
            pic = anc.find('xdr:pic', NS)
            sp = anc.find('xdr:sp', NS)
            if pic is not None:
                blip = pic.find('xdr:blipFill/a:blip', NS)
                rid = blip.get('{%s}embed' % NS['r'])
                crop = pic.find('xdr:blipFill/a:srcRect', NS)
                item = dict(pos, clase='img', medio=medios[rid])
                if crop is not None:
                    item['recorte'] = {k: int(crop.get(k, 0)) / 100000.0 for k in 'ltrb'}
                items.append(item)
            elif sp is not None:
                tx = sp.find('xdr:txBody', NS)
                if tx is None:
                    continue
                bp = tx.find('a:bodyPr', NS)
                fondo = sp.find('xdr:spPr/a:solidFill', NS)
                pars = []
                for pa in tx.findall('a:p', NS):
                    ppr = pa.find('a:pPr', NS)
                    runs = []
                    for rn in pa.findall('a:r', NS):
                        rp = rn.find('a:rPr', NS)
                        run = {'t': rn.find('a:t', NS).text or ''}
                        if rp is not None:
                            if rp.get('sz'): run['sz'] = int(rp.get('sz')) / 100.0
                            if rp.get('b') == '1': run['b'] = 1
                            if rp.get('i') == '1': run['i'] = 1
                            if rp.get('u') and rp.get('u') != 'none': run['u'] = 1
                            c = _color_dml(rp.find('a:solidFill', NS))
                            if c: run['c'] = c
                        runs.append(run)
                    pars.append({'h': (ppr.get('algn') if ppr is not None else None) or 'l', 'runs': runs})
                items.append(dict(pos, clase='texto', parrafos=pars, fondo=_color_dml(fondo) if fondo is not None else None,
                                  ancla=(bp.get('anchor') if bp is not None else None) or 't',
                                  ins=[int(bp.get(k, d)) / EMU for k, d in (('lIns', 91440), ('tIns', 45720), ('rIns', 91440), ('bIns', 45720))] if bp is not None else [9.6, 4.8, 9.6, 4.8]))
        out[nombre] = (items, z)
    return out

def alto_fila(ws, r):
    d = ws.row_dimensions.get(r)
    if d and d.hidden:
        return 0
    return pt_a_px(d.height if d and d.height else (ws.sheet_format.defaultRowHeight or 15))

def ancho_col(ws, ci):
    for dim in ws.column_dimensions.values():
        if dim.min and dim.max and dim.min <= ci <= dim.max:
            return 0 if dim.hidden else px_ancho(dim.width)
    return px_ancho(ws.sheet_format.defaultColWidth)

def main(ruta, salida):
    wb = openpyxl.load_workbook(ruta)
    wv = openpyxl.load_workbook(ruta, data_only=True)
    wr = openpyxl.load_workbook(ruta, rich_text=True)
    cargar_tema(wb)
    dibujos = leer_dibujos(ruta)
    os.makedirs(os.path.join(salida, 'img'), exist_ok=True)
    informe = []
    indice = {}
    imagenes_vistas = {}
    for hoja, clave in HOJAS.items():
        ws, wsv, wsr = wb[hoja], wv[hoja], wr[hoja]
        pa = ws.print_area
        rng = pa.split('!')[1].replace('$', '')
        (c1, r1), (c2, r2) = [(re.match(r'([A-Z]+)(\d+)', p).group(1), int(re.match(r'([A-Z]+)(\d+)', p).group(2))) for p in rng.split(':')]
        ci1, ci2 = column_index_from_string(c1), column_index_from_string(c2)
        cols = []
        for ci in range(ci1, ci2 + 1):
            d = ws.column_dimensions.get(get_column_letter(ci))
            # openpyxl agrupa columnas contiguas con igual ancho: buscar el grupo que contiene a ci
            ancho = None
            oculto = False
            for k, dim in ws.column_dimensions.items():
                if dim.min and dim.max and dim.min <= ci <= dim.max:
                    ancho = dim.width
                    oculto = bool(dim.hidden)
                    break
            cols.append({'px': 0 if oculto else px_ancho(ancho if ancho else ws.sheet_format.defaultColWidth)})
        filas = []
        defecto = ws.sheet_format.defaultRowHeight or 15
        for r in range(r1, r2 + 1):
            d = ws.row_dimensions.get(r)
            oculta = bool(d and d.hidden)
            h = d.height if d and d.height else defecto
            filas.append({'px': 0 if oculta else round(pt_a_px(h), 1)})
        combinadas = {}
        cubiertas = set()
        for mr in ws.merged_cells.ranges:
            if mr.max_row < r1 or mr.min_row > r2 or mr.max_col < ci1 or mr.min_col > ci2:
                continue
            combinadas[(mr.min_row, mr.min_col)] = (min(mr.max_row, r2) - mr.min_row + 1, min(mr.max_col, ci2) - mr.min_col + 1)
            for rr in range(mr.min_row, min(mr.max_row, r2) + 1):
                for cc in range(mr.min_col, min(mr.max_col, ci2) + 1):
                    if (rr, cc) != (mr.min_row, mr.min_col):
                        cubiertas.add((rr, cc))
        celdas = []
        for r in range(r1, r2 + 1):
            for ci in range(ci1, ci2 + 1):
                if (r, ci) in cubiertas:
                    continue
                c = ws.cell(r, ci)
                v = wsv.cell(r, ci).value
                dyn = analizar(c.value, informe, hoja, c.coordinate)
                est = {}
                f = c.font
                if f:
                    if f.b: est['b'] = 1
                    if f.i: est['i'] = 1
                    if f.u: est['u'] = 1
                    if f.sz: est['sz'] = f.sz
                    col = color(f.color)
                    if col and col != '#000000': est['c'] = col
                    if f.name: est['f'] = f.name
                fl = c.fill
                if fl is not None and hasattr(fl, 'stop') and fl.stop:
                    paradas = ', '.join('%s %d%%' % (color(st.color) or '#FFFFFF', round(st.position * 100)) for st in fl.stop)
                    est['bg'] = 'linear-gradient(%ddeg, %s)' % (int(fl.degree or 0) + 90, paradas)
                elif fl is not None and getattr(fl, 'fill_type', None) == 'solid':
                    bg = color(fl.fgColor)
                    if bg and bg.upper() != '#FFFFFF': est['bg'] = bg
                a = c.alignment
                if a:
                    if a.horizontal: est['h'] = a.horizontal
                    if a.vertical: est['v'] = a.vertical
                    if a.wrap_text: est['w'] = 1
                    if a.indent: est['in'] = a.indent
                bd = c.border
                bd_der = bd_aba = bd
                if (r, ci) in combinadas:
                    rs_, cs_ = combinadas[(r, ci)]
                    bd_der = ws.cell(r, ci + cs_ - 1).border
                    bd_aba = ws.cell(r + rs_ - 1, ci).border
                for lado, k, origen in (('top', 't', bd), ('right', 'r', bd_der), ('bottom', 'b', bd_aba), ('left', 'l', bd)):
                    bs = borde(getattr(origen, lado))
                    if bs: est['b' + k] = bs
                hay_texto = v is not None and str(v).strip() != ''
                if not hay_texto and not dyn and not est:
                    continue
                if not hay_texto and not dyn and set(est) <= {'f', 'sz'}:
                    continue
                cel = {'r': r - r1, 'c': ci - ci1}
                if (r, ci) in combinadas:
                    rs, cs = combinadas[(r, ci)]
                    if rs > 1: cel['rs'] = rs
                    if cs > 1: cel['cs'] = cs
                if hay_texto:
                    cel['v'] = v if not hasattr(v, 'isoformat') else v.isoformat()
                if dyn:
                    cel['f'] = dyn
                else:
                    rv = wsr.cell(r, ci).value
                    if isinstance(rv, CellRichText):
                        runs = []
                        for parte in rv:
                            if isinstance(parte, str):
                                runs.append({'t': parte})
                            else:
                                run = {'t': parte.text}
                                ft = parte.font
                                if ft.b: run['b'] = 1
                                if ft.i: run['i'] = 1
                                if ft.sz: run['sz'] = ft.sz
                                cc = color(ft.color)
                                if cc: run['c'] = cc
                                runs.append(run)
                        cel['rt'] = runs
                nf = clase_formato(c.number_format)
                if nf != 'g': cel['nf'] = nf
                if est: cel['s'] = est
                celdas.append(cel)
        # celdas fuera del área de impresión que alimentan fórmulas (columnas auxiliares ocultas)
        aux = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or str(c.value).strip() == '':
                    continue
                if r1 <= c.row <= r2 and ci1 <= c.column <= ci2:
                    continue
                item = {}
                vv = wsv.cell(c.row, c.column).value
                if vv is not None and not hasattr(vv, 'isoformat'): item['v'] = vv
                ar = analizar(c.value, informe, hoja, c.coordinate)
                if ar: item['f'] = ar
                if item: aux[c.coordinate] = item
        imgs = []
        formas = []
        items, zf = dibujos.get(hoja, ([], None))
        for it in items:
            d = it['desde'] if 'desde' in it else None
            base = {}
            if d:
                fr, cr, dx, dy = d['fila'] - (r1 - 1), d['col'] - (ci1 - 1), d['dx'], d['dy']
                if fr < 0:   # arranca sobre el área de impresión: se convierte en desplazamiento negativo
                    dy -= sum(alto_fila(ws, rr) for rr in range(r1 + fr, r1)); fr = 0
                if cr < 0:
                    dx -= sum(ancho_col(ws, cc) for cc in range(ci1 + cr, ci1)); cr = 0
                base = {'fila': fr, 'col': cr, 'dx': round(dx, 1), 'dy': round(dy, 1)}
            if 'hasta' in it:
                h = it['hasta']
                base.update({'fila2': h['fila'] - (r1 - 1), 'col2': h['col'] - (ci1 - 1), 'dx2': round(h['dx'], 1), 'dy2': round(h['dy'], 1)})
            elif 'ext' in it:
                base.update({'ancho': round(it['ext'][0], 1), 'alto': round(it['ext'][1], 1)})
            if it['clase'] == 'img':
                data = zf.read(it['medio'])
                hsh = hashlib.sha1(data).hexdigest()[:12]
                ext = it['medio'].rsplit('.', 1)[-1].lower()
                nombre = '%s.%s' % (hsh, ext)
                ruta_img = os.path.join(salida, 'img', nombre)
                if not os.path.exists(ruta_img):
                    open(ruta_img, 'wb').write(data)
                base['archivo'] = nombre
                if 'recorte' in it: base['recorte'] = it['recorte']
                imgs.append(base)
            else:
                base.update({'parrafos': it['parrafos'], 'ancla': it['ancla'], 'ins': it['ins']})
                if it['fondo']: base['fondo'] = it['fondo']
                formas.append(base)
        cortes = [b.id - (r1 - 1) for b in ws.row_breaks.brk if r1 <= b.id <= r2]
        plantilla = {'clave': clave, 'hoja': hoja, 'columnas': cols, 'filas': filas, 'celdas': celdas,
                     'imagenes': imgs, 'formas': formas, 'cortes': cortes, 'origen': {'fila': r1, 'col': ci1}, 'aux': aux,
                     'escalaExcel': ws.page_setup.scale}
        with open(os.path.join(salida, clave + '.json'), 'w', encoding='utf8') as fh:
            json.dump(plantilla, fh, ensure_ascii=False, separators=(',', ':'))
        indice[clave] = {'hoja': hoja, 'celdas': len(celdas), 'imagenes': len(imgs)}
    with open(os.path.join(salida, 'informe_extraccion.txt'), 'w', encoding='utf8') as fh:
        fh.write('\n'.join(informe) + '\n')
    print(json.dumps(indice, indent=1))
    print('\n'.join(informe))

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
