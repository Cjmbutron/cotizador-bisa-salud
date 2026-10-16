#!/usr/bin/env python3
"""
Extrae las tarifas del cotizador Excel y genera 03_planes_y_tarifas.sql.

Uso:
    pip install openpyxl
    python extraer_tarifas.py COTIZADOR_BISA_SALUD.xlsm 03_planes_y_tarifas.sql \
        --version "Tarifario 2026-05" --desde 2026-05-01

Lee SOLO los datos vigentes confirmados por el negocio (R3):
  - PRIMAS IGR RMA ADV25   B4:BB103   (Infinity, Red Max, Advance por regional)
  - PRIMAS MBL             A6:T87 y V6:AO12 (Mundial Blue adultos e hijos)
  - PRIMAS ALTO RIESGO     (Naranja+, solo ámbitos BOLIVIA y MUNDIAL CON BOLIVIA)
Descarta: PRIMAS IGR RMA ADV25!BE:DE y filas 106-208, PRIMA FILTRADO, DATOS LISTAS A-L.

La versión se crea en estado BORRADOR. Se publica después de la conciliación
(ver README), con un aprobador distinto del creador.
"""
import argparse
import sys
from decimal import Decimal
import openpyxl
from openpyxl.utils import column_index_from_string as ci

REGIONES = {'LPZ': 'B', 'CBBA': 'M', 'SCZ': 'X', 'CHQ': 'AI', 'TJA': 'AT'}
# columnas dentro de cada bloque regional: producto -> (offset, plan) ; sexo en orden M, F-M, F
IGR_PRODUCTOS = {'IG': (0, 'IG'), 'RM': (3, 'RM_700K'), 'AD': (6, 'AD')}
SEXOS = ['M', 'F-M', 'F']

COBERTURAS = {'1MM': (0, 1_000_000), '2MM': (6, 2_000_000), '5MM': (12, 5_000_000)}
DEDUCIBLES_MB = {'1K': (0, 1000), '2K': (1, 2000), '3K': (2, 3000), '5K': (3, 5000), '10K': (4, 10000), '20K': (5, 20000)}

NP_VA = {'50K': (0, 50_000), '100K': (1, 100_000), '150K': (2, 150_000), '250K': (3, 250_000),
         '500K': (4, 500_000), '1MM': (5, 1_000_000), '2MM': (6, 2_000_000)}
NP_DED = {'2K': (0, 2000), '3K': (1, 3000), '4K': (2, 4000), '5K': (3, 5000), '10K': (4, 10000), '15K': (5, 15000),
          '20K': (6, 20000), '25K': (7, 25000), '30K': (8, 30000), '40K': (9, 40000), '50K': (10, 50000), '100K': (11, 100000)}
# ámbito -> (código corto, valores asegurados y deducibles ofrecidos en las listas del Excel)
NP_AMBITOS = {
    'BOLIVIA':             ('L', ['50K', '100K', '250K'],                 ['5K', '10K']),
    'MUNDIAL_CON_BOLIVIA': ('MB', ['150K', '250K', '500K', '1MM'],        ['5K', '10K', '50K', '100K']),
}


def num(v):
    """Número o None. El texto 'N/A', celdas vacías y cualquier texto -> None (no disponible)."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    return None


def fmt(v):
    if v is None:
        return 'NULL'
    d = Decimal(str(v)).quantize(Decimal('0.000001'))
    s = format(d, 'f').rstrip('0').rstrip('.')
    return s if s else '0'


def lit(s):
    return "'" + str(s).replace("'", "''") + "'"


def batches(rows, n=1000):
    for i in range(0, len(rows), n):
        yield rows[i:i + n]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsm'); ap.add_argument('salida')
    ap.add_argument('--version', default='Tarifario 2026-05')
    ap.add_argument('--desde', default='2026-05-01')
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsm, data_only=True)
    igr, mbl, ar = wb['PRIMAS IGR RMA ADV25'], wb['PRIMAS MBL'], wb['PRIMAS ALTO RIESGO']
    out = []
    w = out.append

    w('-- Generado por extraer_tarifas.py. No editar a mano.')
    w('BEGIN;')
    w(f"INSERT INTO version_tarifa (nombre, vigente_desde, estado, respaldo_documental) VALUES "
      f"({lit(a.version)}, DATE {lit(a.desde)}, 'BORRADOR', {lit('Importado de ' + a.xlsm.split('/')[-1])});")

    # ---------------- planes ----------------
    planes = []   # (producto, codigo, nombre, ambito, va, moneda_va, deducible)
    planes.append(('IG', 'IG', 'Infinity Green', None, None, None, None))
    planes.append(('RM', 'RM_700K', 'Red Max Bs. 700.000', None, 700000, 'BOB', None))
    planes.append(('RM', 'RM_140K', 'Red Max Bs. 140.000', None, 140000, 'BOB', None))
    planes.append(('RM', 'RM_70K', 'Red Max Bs. 70.000', None, 70000, 'BOB', None))
    planes.append(('AD', 'AD', 'Advance', None, None, None, None))
    for cov, (_, va) in COBERTURAS.items():
        for ded, (_, dv) in DEDUCIBLES_MB.items():
            planes.append(('MB', f'MB_{cov}_{ded}', f'Mundial Blue {cov} deducible {ded}', None, va, 'USD', dv))
    for amb, (cod, vas, deds) in NP_AMBITOS.items():
        for va in vas:
            for ded in deds:
                planes.append(('NP', f'NP_{cod}_{va}_{ded}', f'Naranja+ {amb.replace("_", " ").title()} {va} deducible {ded}',
                               amb, NP_VA[va][1], 'USD', NP_DED[ded][1]))
    vals = ',\n  '.join('(' + ', '.join('NULL' if x is None else (lit(x) if isinstance(x, str) else str(x)) for x in p) + ')' for p in planes)
    w("INSERT INTO plan (producto_id, codigo, nombre, ambito, valor_asegurado, moneda_va, deducible)\n"
      "SELECT pr.id, v.codigo, v.nombre, v.ambito, v.va::numeric, v.moneda, v.ded::numeric\n"
      f"FROM (VALUES\n  {vals}\n) AS v(prod, codigo, nombre, ambito, va, moneda, ded)\nJOIN producto pr ON pr.codigo = v.prod;")
    w("UPDATE plan SET plan_base_id = (SELECT id FROM plan WHERE codigo = 'RM_700K') WHERE codigo IN ('RM_140K','RM_70K');")

    filas_tarifa = []   # (plan, region, sexo, edad, prima, periodicidad)
    # ---------------- Infinity / Red Max / Advance ----------------
    for reg, col in REGIONES.items():
        c0 = ci(col)
        for prod, (off, plan) in IGR_PRODUCTOS.items():
            for k, sexo in enumerate(SEXOS):
                for edad in range(0, 100):
                    filas_tarifa.append((plan, reg, sexo, edad, num(igr.cell(4 + edad, c0 + off + k).value), 'MENSUAL'))
    n_igr = len(filas_tarifa)
    # ---------------- Mundial Blue adultos ----------------
    for cov, (co, _) in COBERTURAS.items():
        for ded, (do, _) in DEDUCIBLES_MB.items():
            for edad in range(18, 100):
                filas_tarifa.append((f'MB_{cov}_{ded}', None, None, edad, num(mbl.cell(6 + edad - 18, 3 + co + do).value), 'MENSUAL'))
    n_mb = len(filas_tarifa) - n_igr
    # ---------------- Naranja+ (prima ANUAL, edad 0..84) ----------------
    for amb, (cod, vas, deds) in NP_AMBITOS.items():
        for va in vas:
            for ded in deds:
                col = 2 + 12 * NP_VA[va][0] + 84 * {'L': 0, 'U': 1, 'M': 2, 'MB': 3}[cod] + NP_DED[ded][0]
                for edad in range(0, 85):
                    v = ar.cell(18 + edad, col).value
                    if v is None:
                        continue                      # celda vacía: la combinación no existe a esa edad
                    filas_tarifa.append((f'NP_{cod}_{va}_{ded}', None, None, edad, num(v), 'ANUAL'))
    n_np = len(filas_tarifa) - n_igr - n_mb

    for lote in batches(filas_tarifa):
        vals = ',\n  '.join(f"({lit(p)}, {lit(r) if r else 'NULL'}, {lit(s) if s else 'NULL'}, {e}, {fmt(pr)}, {lit(per)})"
                            for p, r, s, e, pr, per in lote)
        w("INSERT INTO tarifa (version_id, plan_id, region_id, sexo, edad_desde, edad_hasta, prima, periodicidad, disponible)\n"
          f"SELECT (SELECT id FROM version_tarifa WHERE nombre = {lit(a.version)}), p.id, r.id, v.sexo, v.edad, v.edad,\n"
          "       v.prima::numeric, v.per::periodicidad, v.prima IS NOT NULL\n"
          f"FROM (VALUES\n  {vals}\n) AS v(plan, region, sexo, edad, prima, per)\n"
          "JOIN plan p ON p.codigo = v.plan\nLEFT JOIN region r ON r.codigo = v.region;")

    # ---------------- Mundial Blue hijos (prima POR hijo, ya dividida) ----------------
    clave_fila = {}
    for r in range(6, 13):
        n = mbl.cell(r, 22).value
        if isinstance(n, (int, float)):
            clave_fila[int(n)] = r
    hijos = []
    for cov, (co, _) in COBERTURAS.items():
        for ded, (do, _) in DEDUCIBLES_MB.items():
            for n in range(1, 6):
                v = num(mbl.cell(clave_fila[n], 24 + co + do).value)
                if v is None:
                    sys.exit(f'Sin tarifa de hijo para {cov} {ded} con {n} hijos')
                hijos.append((f'MB_{cov}_{ded}', n, v))
    vals = ',\n  '.join(f"({lit(p)}, {n}, {fmt(v)})" for p, n, v in hijos)
    w("INSERT INTO tarifa_hijo (version_id, plan_id, numero_hijos, prima_por_hijo)\n"
      f"SELECT (SELECT id FROM version_tarifa WHERE nombre = {lit(a.version)}), p.id, v.n, v.prima\n"
      f"FROM (VALUES\n  {vals}\n) AS v(plan, n, prima)\nJOIN plan p ON p.codigo = v.plan;")
    w('COMMIT;')

    with open(a.salida, 'w', encoding='utf8') as f:
        f.write('\n'.join(out) + '\n')
    print(f'Planes: {len(planes)} | tarifas IG/RM/AD: {n_igr} | MB adultos: {n_mb} | Naranja+: {n_np} | MB hijos: {len(hijos)}')


if __name__ == '__main__':
    main()
