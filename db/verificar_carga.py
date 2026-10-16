#!/usr/bin/env python3
"""
Concilia, fila por fila, la tarifa cargada en PostgreSQL contra el Excel original.

Uso:
    pip install openpyxl psycopg2-binary
    python verificar_carga.py COTIZADOR_BISA_SALUD.xlsm "postgresql://usuario@host/cotizador" "Tarifario 2026-05"

Sale con código 1 si encuentra cualquier diferencia. Pensado para ejecutarse antes de
publicar cada versión de tarifas (versión nueva contra su Excel de origen).
"""
import sys
from decimal import Decimal
import openpyxl
import psycopg2
from openpyxl.utils import column_index_from_string as ci

REG = {'LPZ': 'B', 'CBBA': 'M', 'SCZ': 'X', 'CHQ': 'AI', 'TJA': 'AT'}
COV = {'1MM': 0, '2MM': 6, '5MM': 12}
DED = {'1K': 0, '2K': 1, '3K': 2, '5K': 3, '10K': 4, '20K': 5}
VA = {'50K': 0, '100K': 1, '150K': 2, '250K': 3, '500K': 4, '1MM': 5}
DD = {'5K': 3, '10K': 4, '50K': 10, '100K': 11}
NP = {'L': (0, ['50K', '100K', '250K'], ['5K', '10K']), 'MB': (3, ['150K', '250K', '500K', '1MM'], ['5K', '10K', '50K', '100K'])}


def n(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def main(xlsm, dsn, version):
    wb = openpyxl.load_workbook(xlsm, data_only=True)
    igr, mbl, ar = wb['PRIMAS IGR RMA ADV25'], wb['PRIMAS MBL'], wb['PRIMAS ALTO RIESGO']
    cur = psycopg2.connect(dsn).cursor()
    cur.execute("""SELECT p.codigo, r.codigo, t.sexo, t.edad_desde, t.prima, t.disponible
                     FROM tarifa t JOIN version_tarifa v ON v.id = t.version_id
                     JOIN plan p ON p.id = t.plan_id LEFT JOIN region r ON r.id = t.region_id
                    WHERE v.nombre = %s""", (version,))
    db = {(p, r, s, e): (None if pr is None else Decimal(pr), d) for p, r, s, e, pr, d in cur.fetchall()}
    esperado = {}
    for reg, col in REG.items():
        c0 = ci(col)
        for plan, off in (('IG', 0), ('RM_700K', 3), ('AD', 6)):
            for k, sx in enumerate(['M', 'F-M', 'F']):
                for e in range(100):
                    esperado[(plan, reg, sx, e)] = n(igr.cell(4 + e, c0 + off + k).value)
    for cov, co in COV.items():
        for ded, do in DED.items():
            for e in range(18, 100):
                esperado[(f'MB_{cov}_{ded}', None, None, e)] = n(mbl.cell(6 + e - 18, 3 + co + do).value)
    for cod, (sc, vas, deds) in NP.items():
        for va in vas:
            for ded in deds:
                for e in range(85):
                    v = ar.cell(18 + e, 2 + 12 * VA[va] + 84 * sc + DD[ded]).value
                    if v is not None:
                        esperado[(f'NP_{cod}_{va}_{ded}', None, None, e)] = n(v)
    errores = []
    if set(db) != set(esperado):
        errores.append(f'Claves distintas: solo en BD {len(set(db)-set(esperado))}, solo en Excel {len(set(esperado)-set(db))}')
    for k, ev in esperado.items():
        if k not in db:
            continue
        dv, disp = db[k]
        if ev is None:
            ok = dv is None and not disp
        else:
            ok = dv is not None and abs(dv - Decimal(str(ev))) < Decimal('0.000001') and disp
        if not ok:
            errores.append(f'{k}: Excel={ev} BD={dv}')
    # tarifas de hijo
    cur.execute("""SELECT p.codigo, h.numero_hijos, h.prima_por_hijo FROM tarifa_hijo h
                     JOIN version_tarifa v ON v.id = h.version_id JOIN plan p ON p.id = h.plan_id WHERE v.nombre = %s""", (version,))
    hdb = {(p, k): Decimal(v) for p, k, v in cur.fetchall()}
    filas = {int(mbl.cell(r, 22).value): r for r in range(6, 13) if isinstance(mbl.cell(r, 22).value, (int, float))}
    nh = 0
    for cov, co in COV.items():
        for ded, do in DED.items():
            for k in range(1, 6):
                ev = Decimal(str(mbl.cell(filas[k], 24 + co + do).value))
                nh += 1
                if (f'MB_{cov}_{ded}', k) not in hdb or abs(hdb[(f'MB_{cov}_{ded}', k)] - ev) >= Decimal('0.000001'):
                    errores.append(f'hijo {cov} {ded} {k}: Excel={ev} BD={hdb.get((f"MB_{cov}_{ded}", k))}')
    print(f'Comparadas {len(esperado)} tarifas y {nh} tarifas de hijo.')
    if errores:
        print(f'{len(errores)} DIFERENCIAS:')
        print('\n'.join(errores[:30]))
        sys.exit(1)
    print('Conciliación exacta: 0 diferencias.')


if __name__ == '__main__':
    main(*sys.argv[1:4])
